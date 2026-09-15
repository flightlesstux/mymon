"""World Bank Commodity Price Data ("Pink Sheet"), monthly XLSX.

The document id in the XLSX URL changes with every release (the ``-0350012021`` file listed
in older docs still downloads but stops at 2024M12), so ``fetch`` first scrapes the current
``CMO-Historical-Data-Monthly.xlsx`` link from the commodity-markets page and only falls
back to the fixed URL when the page is unreachable or has no link.

``Monthly Prices`` sheet layout (verified live):

- a few title rows, then a names row (``Crude oil, Brent``, ``Natural gas, Europe`` ...,
  first cell empty), then a units row (``($/bbl)``, ``($/mmbtu)`` ...),
- data rows whose first cell is ``YYYYMmm`` (``1960M01``); missing values are ``…``.

The names/units rows are located relative to the first ``YYYYMmm`` row, so extra title
rows do not matter. Commodity keys are slugs of the header (``crude_oil_brent``, ``gold``);
the unit is the units-row text without the parentheses (``$/bbl``). The file holds the full
history, so there is no separate backfill.
"""

from __future__ import annotations

import html
import io
import logging
import re
from datetime import date
from typing import Any

import openpyxl

from ..source import Ctx, Rows, Source

log = logging.getLogger(__name__)

PAGE_URL = "https://www.worldbank.org/en/research/commodity-markets"
FALLBACK_URL = (
    "https://thedocs.worldbank.org/en/doc/5d903e848db1d1b83e0ec8f744e55570-0350012021"
    "/related/CMO-Historical-Data-Monthly.xlsx"
)
SOURCE_NAME = "wb_pink"
SHEET_NAME = "Monthly Prices"

_LINK_RE = re.compile(
    r"""https?://[^"'\s<>]+/CMO-Historical-Data-Monthly\.xlsx""", re.IGNORECASE
)
_PERIOD_RE = re.compile(r"^\s*(\d{4})M(\d{1,2})\s*$")
_MISSING = {"", "…", "...", "-", "n/a", "na"}


# --------------------------------------------------------------------------- helpers


def discover_xlsx_url(page_html: str) -> str | None:
    m = _LINK_RE.search(page_html)
    return html.unescape(m.group(0)) if m else None


def slugify(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
    return slug


def _clean_unit(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    unit = value.strip()
    if unit.startswith("(") and unit.endswith(")"):
        unit = unit[1:-1].strip()
    return unit or None


def _num(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        value = value.strip()
        if value.lower() in _MISSING:
            return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _period(value: Any) -> date | None:
    if not isinstance(value, str):
        return None
    m = _PERIOD_RE.match(value)
    if not m:
        return None
    year, month = int(m.group(1)), int(m.group(2))
    if not 1 <= month <= 12:
        return None
    return date(year, month, 1)


def _pick_sheet(names: list[str]) -> str | None:
    if SHEET_NAME in names:
        return SHEET_NAME
    for name in names:
        if "monthly" in name.lower() and "price" in name.lower():
            return name
    return None


def _names_row(buffer: list[tuple[Any, ...]]) -> tuple[Any, ...] | None:
    """Latest buffered row (above the units row) with at least three text cells."""
    for row in reversed(buffer):
        texts = [c for c in row[1:] if isinstance(c, str) and c.strip()]
        if len(texts) >= 3 and not all(t.strip().startswith("(") for t in texts):
            return row
    return None


def parse_workbook(content: bytes) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = _pick_sheet(wb.sheetnames)
        if sheet is None:
            raise ValueError(f"no '{SHEET_NAME}' sheet in {wb.sheetnames}")
        ws = wb[sheet]
        buffer: list[tuple[Any, ...]] = []
        columns: dict[int, tuple[str, str | None]] = {}
        rows: list[dict[str, Any]] = []
        for row in ws.iter_rows(values_only=True):
            period = _period(row[0] if row else None)
            if period is None:
                if not columns:
                    buffer.append(row)
                    buffer = buffer[-8:]
                continue
            if not columns:
                units = buffer[-1] if buffer else ()
                names = _names_row(buffer[:-1])
                if names is None:
                    raise ValueError("no commodity names row above the first data row")
                for idx, cell in enumerate(names):
                    if idx == 0 or not isinstance(cell, str) or not cell.strip():
                        continue
                    slug = slugify(cell)
                    if not slug:
                        continue
                    unit = _clean_unit(units[idx]) if idx < len(units) else None
                    columns[idx] = (slug, unit)
                if not columns:
                    raise ValueError("commodity names row is empty")
            for idx, (slug, unit) in columns.items():
                value = _num(row[idx]) if idx < len(row) else None
                if value is None:
                    continue
                rows.append(
                    {
                        "period_date": period,
                        "commodity": slug,
                        "price": value,
                        "unit": unit,
                        "source": SOURCE_NAME,
                    }
                )
        if not columns:
            raise ValueError("no YYYYMmm data rows found")
        return rows
    finally:
        wb.close()


# --------------------------------------------------------------------------- source


def _resolve_url(ctx: Ctx) -> str:
    try:
        page = ctx.http.get(PAGE_URL)
        page.raise_for_status()
    except Exception as exc:  # noqa: BLE001 - any page failure falls back to the fixed URL
        log.warning("commodities_wb: page fetch failed (%s), using fallback URL", exc)
        return FALLBACK_URL
    url = discover_xlsx_url(page.text)
    if url is None:
        log.warning("commodities_wb: no monthly XLSX link on page, using fallback URL")
        return FALLBACK_URL
    return url


def fetch(ctx: Ctx) -> Rows:
    url = _resolve_url(ctx)
    log.info("commodities_wb: downloading %s", url)
    resp = ctx.http.get(url)
    resp.raise_for_status()
    rows = parse_workbook(resp.content)
    log.info("commodities_wb: %d rows", len(rows))
    return [("commodity_price", rows)]


SOURCE = Source(
    name="commodities_wb",
    interval=86400,
    fetch=fetch,
    backfill=None,
    tables=["commodity_price"],
    description="World Bank Pink Sheet monthly commodity prices (full history).",
)
