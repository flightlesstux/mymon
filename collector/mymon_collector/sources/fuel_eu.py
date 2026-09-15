"""EU fuel prices: European Commission Weekly Oil Bulletin price history.

The bulletin page links a single XLSX (``Weekly_Oil_Bulletin_Prices_History_*.xlsx``) that
carries the whole weekly history since 2005, so ``fetch`` re-downloads it every day and
returns the full history; there is no separate backfill.

Layout of the ``Prices with taxes`` sheet (verified against the live file):

- row 1: machine ids, one 7-column block per country: ``CTR``, ``AT_price_with_tax_euro95``,
  ``AT_price_with_tax_diesel``, ``AT_price_with_tax_heating_oil``, ``..fuel_oil_1``,
  ``..fuel_oil_2``, ``AT_price_with_tax_LPG``; blocks for ``EU_`` (EU average), ``EUR_``
  (euro-area average) and every member state, plus a legacy ``UK_`` block.
- row 2: human labels (``Euro-super 95 (I)``, ``Gas oil automobile ...``), row 3: units
  (``1000 l``), row 4 onwards: one row per Monday, newest first, first cell a datetime.
- trailing rows with an empty date cell hold footnotes.

Columns are located purely by the row-1 ids, so column order and extra blocks do not matter.
Only two-letter codes are emitted (``EU`` is kept as the EU-wide aggregate, ``EUR`` is
dropped), prices are converted from EUR per 1000 l to EUR per litre.
"""

from __future__ import annotations

import html
import io
import logging
import re
from datetime import date, datetime
from typing import Any
from urllib.parse import urljoin

import openpyxl

from ..source import Ctx, Rows, Source

log = logging.getLogger(__name__)

PAGE_URL = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
SOURCE_NAME = "eu_wob"

FUEL_TYPES = {"euro95": "petrol95", "diesel": "diesel", "lpg": "lpg"}
_ID_RE = re.compile(r"^([A-Za-z]{2,3})_price_with_tax_(euro95|diesel|LPG)$", re.IGNORECASE)
_HREF_RE = re.compile(r"""href=["']([^"']+\.xlsx[^"']*)["']""", re.IGNORECASE)
_LINK_PRIORITY = ("prices_history", "price_history", "history")


# --------------------------------------------------------------------------- helpers


def discover_xlsx_url(page_html: str, base_url: str = PAGE_URL) -> str | None:
    """Pick the price-history XLSX link out of the bulletin page."""
    links = [html.unescape(h) for h in _HREF_RE.findall(page_html)]
    for needle in _LINK_PRIORITY:
        for link in links:
            if needle in link.lower():
                return urljoin(base_url, link)
    return None


def _pick_sheet(names: list[str]) -> str | None:
    for name in names:
        low = name.lower()
        if "with tax" in low and "wo" not in low.split():
            return name
    for name in names:
        if "with" in name.lower() and "tax" in name.lower():
            return name
    return None


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip()[:10]).date()
        except ValueError:
            return None
    return None


def _num(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        value = value.strip().replace(",", ".")
        if not value:
            return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _column_map(row: tuple[Any, ...]) -> dict[int, tuple[str, str]]:
    """``{column_index: (country_iso2, fuel_type)}`` from the machine-id row."""
    out: dict[int, tuple[str, str]] = {}
    for idx, cell in enumerate(row):
        if not isinstance(cell, str):
            continue
        m = _ID_RE.match(cell.strip())
        if not m:
            continue
        code = m.group(1).upper()
        if len(code) != 2:
            continue  # EUR_ euro-area aggregate
        out[idx] = (code, FUEL_TYPES[m.group(2).lower()])
    return out


def parse_workbook(content: bytes) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet_name = _pick_sheet(wb.sheetnames)
        if sheet_name is None:
            raise ValueError(f"no 'prices with taxes' sheet in {wb.sheetnames}")
        ws = wb[sheet_name]
        columns: dict[int, tuple[str, str]] = {}
        rows: list[dict[str, Any]] = []
        skipped = 0
        for row in ws.iter_rows(values_only=True):
            if not columns:
                columns = _column_map(row)
                continue
            day = _as_date(row[0] if row else None)
            if day is None:
                continue  # label/unit rows above and footnotes below the data
            for idx, (country, fuel) in columns.items():
                value = _num(row[idx]) if idx < len(row) else None
                if value is None:
                    continue
                if value <= 0:
                    skipped += 1
                    continue
                rows.append(
                    {
                        "period_date": day,
                        "country_iso2": country,
                        "region": "",
                        "fuel_type": fuel,
                        "price": round(value / 1000.0, 4),
                        "currency": "EUR",
                        "unit": "EUR/L",
                        "source": SOURCE_NAME,
                    }
                )
        if not columns:
            raise ValueError("no *_price_with_tax_* id row found")
        if skipped:
            log.warning("fuel_eu: skipped %d non-positive prices", skipped)
        return rows
    finally:
        wb.close()


# --------------------------------------------------------------------------- source


def fetch(ctx: Ctx) -> Rows:
    page = ctx.http.get(PAGE_URL)
    page.raise_for_status()
    url = discover_xlsx_url(page.text)
    if url is None:
        raise ValueError("fuel_eu: no price-history XLSX link on the bulletin page")
    log.info("fuel_eu: downloading %s", url)
    resp = ctx.http.get(url)
    resp.raise_for_status()
    rows = parse_workbook(resp.content)
    log.info("fuel_eu: %d rows", len(rows))
    return [("fuel_price", rows)]


SOURCE = Source(
    name="fuel_eu",
    interval=86400,
    fetch=fetch,
    backfill=None,
    tables=["fuel_price"],
    description="EU Weekly Oil Bulletin pump prices incl. taxes (EUR/L, full weekly history).",
)
